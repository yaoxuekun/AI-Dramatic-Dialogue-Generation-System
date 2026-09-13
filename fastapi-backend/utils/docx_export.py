"""漫剧剧本 Word (.docx) 导出工具。"""

from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from repositories import story_repository
from schemas.story import (
    CharacterResponse,
    VolumeOutlineResponse,
    VolumeSectionResponse,
    SectionScriptResponse,
    AssetResponse,
)
import database


def export_story_docx(story_id: int) -> BytesIO:
    """查询故事全部数据，生成 Word 文档并返回 BytesIO。"""
    with database.get_db_cursor() as conn:
        story = story_repository.find_by_id(conn, story_id)
        if not story:
            raise ValueError("故事不存在")

        characters_raw = story_repository.find_characters_by_story_id(conn, story_id)
        characters = [CharacterResponse(**c) for c in characters_raw]

        volumes_raw = story_repository.find_volume_outlines_by_story_id(conn, story_id)
        volume_ids = [v["id"] for v in volumes_raw]

        sections_by_volume = {}
        section_ids = []
        if volume_ids:
            sections_raw = story_repository.find_sections_by_volume_ids(conn, volume_ids)
            for s in sections_raw:
                vid = s["volume_id"]
                sections_by_volume.setdefault(vid, []).append(s)
                section_ids.append(s["id"])

        scripts_by_section = {}
        if section_ids:
            scripts_raw = story_repository.find_scripts_by_section_ids(conn, section_ids)
            for sc in scripts_raw:
                sid = sc["section_id"]
                scripts_by_section.setdefault(sid, []).append(sc)

    # ── 构建 Word 文档 ──────────────────────────────────────
    doc = Document()

    # 标题
    title_para = doc.add_heading(level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.add_run(story["title"] or "漫剧剧本")

    # 一、故事摘要
    doc.add_heading("一、故事摘要", level=1)
    doc.add_paragraph(story.get("synopsis") or "暂无摘要。")

    # 二、剧情大纲
    doc.add_heading("二、剧情大纲", level=1)
    doc.add_paragraph(story.get("full_content") or "暂无剧情大纲。")

    # 三、主要角色
    if characters:
        doc.add_heading("三、主要角色", level=1)
        for i, c in enumerate(characters, 1):
            doc.add_heading(f"3.{i} {c.name} — {c.role or ''}", level=2)
            if c.description:
                doc.add_paragraph(f"描述：{c.description}")
            if c.personality:
                doc.add_paragraph(f"性格：{c.personality}")

    # 四、分卷大纲
    if volumes_raw:
        doc.add_heading("四、分卷大纲", level=1)
        for v in volumes_raw:
            doc.add_heading(f"第{v['volume_number']}卷：{v['title']}", level=2)
            if v.get("summary"):
                doc.add_paragraph(v["summary"])

    # 五、分卷详情（含小节内容和分镜脚本）
    if volumes_raw:
        doc.add_heading("五、分卷详情", level=1)
        for v in volumes_raw:
            vn = v["volume_number"]
            doc.add_heading(f"第{vn}卷：{v['title']}", level=2)

            sections = sections_by_volume.get(v["id"], [])
            for s in sections:
                doc.add_heading(f"{s['section_number']}. {s['title']}", level=3)
                if s.get("content"):
                    doc.add_paragraph(s["content"])

                # 分镜脚本
                scripts = scripts_by_section.get(s["id"], [])
                if scripts:
                    doc.add_heading("分镜脚本", level=4)
                    # 表格：镜头号 | 类型 | 时长 | 动作 | 台词
                    table = doc.add_table(rows=1, cols=5, style="Light Grid Accent 1")
                    hdr = table.rows[0].cells
                    hdr[0].text = "镜头"
                    hdr[1].text = "类型"
                    hdr[2].text = "时长(秒)"
                    hdr[3].text = "动作"
                    hdr[4].text = "台词"
                    for sc in sorted(scripts, key=lambda x: x.get("shot_number", 0)):
                        row = table.add_row().cells
                        row[0].text = str(sc.get("shot_number", ""))
                        row[1].text = sc.get("shot_type", "")
                        row[2].text = str(sc.get("duration_seconds", ""))
                        row[3].text = sc.get("action", "")
                        row[4].text = sc.get("dialogue", "") or ""

    # 输出到 BytesIO
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
